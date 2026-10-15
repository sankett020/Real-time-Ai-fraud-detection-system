import sys
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Color Palette - Sentinel AI Cyber / FinTech Navy Theme
    BG_DARK = RGBColor(11, 19, 43)        # #0B132B - Deep Navy
    CARD_BG = RGBColor(19, 34, 64)        # #132240 - Dark Slate Card
    CARD_BORDER = RGBColor(28, 52, 94)    # #1C345E - Subtle Card Border
    ACCENT_CYAN = RGBColor(0, 229, 255)   # #00E5FF - Vibrant Cyber Cyan
    ACCENT_TEAL = RGBColor(79, 172, 254)  # #4FACFE - Electric Blue
    ACCENT_GREEN = RGBColor(0, 230, 118)  # #00E676 - Bright Green
    ACCENT_RED = RGBColor(255, 82, 82)    # #FF5252 - Warning Red
    ACCENT_AMBER = RGBColor(255, 177, 66) # #FFB142 - Alert Amber
    TEXT_WHITE = RGBColor(240, 244, 248)  # #F0F4F8 - Pure Light
    TEXT_MUTED = RGBColor(156, 175, 198)  # #9CAFC6 - Soft Subtext
    TEXT_DIM = RGBColor(100, 120, 145)    # #647891 - Darker muted

    def set_slide_background(slide):
        bg_shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
        bg_shape.fill.solid()
        bg_shape.fill.fore_color.rgb = BG_DARK
        bg_shape.line.color.rgb = BG_DARK
        return bg_shape

    def add_header(slide, title_text, category_text="SENTINEL AI // SYSTEM OVERVIEW", slide_num=None):
        set_slide_background(slide)
        
        # Category Tag
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(9.0), Inches(0.4))
        tf_cat = cat_box.text_frame
        tf_cat.word_wrap = True
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.name = "Arial"
        p_cat.font.size = Pt(11)
        p_cat.font.bold = True
        p_cat.font.color.rgb = ACCENT_CYAN

        # Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(10.5), Inches(0.8))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.name = "Arial"
        p_title.font.size = Pt(24)
        p_title.font.bold = True
        p_title.font.color.rgb = TEXT_WHITE

        # Divider line
        div = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.5), Inches(11.733), Inches(0.03))
        div.fill.solid()
        div.fill.fore_color.rgb = CARD_BORDER
        div.line.color.rgb = CARD_BORDER

        # Slide Number
        if slide_num:
            num_box = slide.shapes.add_textbox(Inches(11.0), Inches(0.5), Inches(1.5), Inches(0.4))
            p_num = num_box.text_frame.paragraphs[0]
            p_num.alignment = PP_ALIGN.RIGHT
            p_num.text = f"{slide_num:02d} / 13"
            p_num.font.name = "Arial"
            p_num.font.size = Pt(11)
            p_num.font.color.rgb = TEXT_DIM

    def add_card(slide, left, top, width, height, title=None, border_color=CARD_BORDER, bg_color=CARD_BG):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        card.line.color.rgb = border_color
        card.line.width = Pt(1.2)
        
        if title:
            tb = slide.shapes.add_textbox(left + Inches(0.2), top + Inches(0.15), width - Inches(0.4), Inches(0.45))
            p = tb.text_frame.paragraphs[0]
            p.text = title
            p.font.name = "Arial"
            p.font.size = Pt(13)
            p.font.bold = True
            p.font.color.rgb = ACCENT_CYAN
        return card

    # =========================================================================
    # SLIDE 1: TITLE & TEAM PRESENTATION
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1)

    # Accent Top Banner / Brand Indicator
    badge = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.9), Inches(3.2), Inches(0.4))
    badge.fill.solid()
    badge.fill.fore_color.rgb = CARD_BG
    badge.line.color.rgb = ACCENT_CYAN
    badge.line.width = Pt(1)
    p_b = badge.text_frame.paragraphs[0]
    p_b.text = "INTELLIGENT CYBERSECURITY PLATFORM"
    p_b.font.size = Pt(10)
    p_b.font.bold = True
    p_b.font.color.rgb = ACCENT_CYAN
    p_b.alignment = PP_ALIGN.CENTER

    # Project Title
    t_box = s1.shapes.add_textbox(Inches(0.8), Inches(1.4), Inches(11.5), Inches(1.8))
    tf1 = t_box.text_frame
    p1 = tf1.paragraphs[0]
    p1.text = "SENTINEL AI"
    p1.font.name = "Arial Black"
    p1.font.size = Pt(40)
    p1.font.bold = True
    p1.font.color.rgb = TEXT_WHITE

    p1_sub = tf1.add_paragraph()
    p1_sub.text = "Real-Time Intelligent Financial Fraud Detection System"
    p1_sub.font.name = "Arial"
    p1_sub.font.size = Pt(20)
    p1_sub.font.color.rgb = ACCENT_TEAL
    p1_sub.space_before = Pt(8)

    p1_desc = tf1.add_paragraph()
    p1_desc.text = "An end-to-end full-stack platform integrating machine learning inference, explainable behavioral risk scoring, transaction monitoring, and security investigation workflows."
    p1_desc.font.name = "Arial"
    p1_desc.font.size = Pt(13)
    p1_desc.font.color.rgb = TEXT_MUTED
    p1_desc.space_before = Pt(10)

    # Team Card
    add_card(s1, Inches(0.8), Inches(4.3), Inches(11.733), Inches(2.4), "PROJECT TEAM MEMBERS", border_color=ACCENT_CYAN)

    # Team Leader Box
    tl_card = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.1), Inches(4.9), Inches(3.6), Inches(1.5))
    tl_card.fill.solid()
    tl_card.fill.fore_color.rgb = RGBColor(14, 25, 50)
    tl_card.line.color.rgb = ACCENT_CYAN
    tl_card.line.width = Pt(1.5)
    tf_tl = tl_card.text_frame
    tf_tl.vertical_anchor = MSO_ANCHOR.MIDDLE
    p_tl_role = tf_tl.paragraphs[0]
    p_tl_role.text = "★ TEAM LEADER"
    p_tl_role.font.size = Pt(11)
    p_tl_role.font.bold = True
    p_tl_role.font.color.rgb = ACCENT_CYAN
    p_tl_name = tf_tl.add_paragraph()
    p_tl_name.text = "Sanket Kushwaha"
    p_tl_name.font.size = Pt(16)
    p_tl_name.font.bold = True
    p_tl_name.font.color.rgb = TEXT_WHITE
    p_tl_name.space_before = Pt(4)
    p_tl_desc = tf_tl.add_paragraph()
    p_tl_desc.text = "Full-Stack & ML Architecture"
    p_tl_desc.font.size = Pt(11)
    p_tl_desc.font.color.rgb = TEXT_MUTED

    # Team Members (3 members)
    members = [
        ("Anand Verma", "Backend & API Engineering", "FastAPI, DB & Services"),
        ("Sandeep Mishra", "Machine Learning Pipeline", "Model Training & Evaluation"),
        ("Kunal", "Frontend UI/UX Engineering", "React, Tailwind & Recharts")
    ]
    for i, (name, role, desc) in enumerate(members):
        m_card = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(4.9 + i * 2.5), Inches(4.9), Inches(2.35), Inches(1.5))
        m_card.fill.solid()
        m_card.fill.fore_color.rgb = RGBColor(14, 25, 50)
        m_card.line.color.rgb = CARD_BORDER
        m_card.line.width = Pt(1)
        tf_m = m_card.text_frame
        tf_m.vertical_anchor = MSO_ANCHOR.MIDDLE
        p_m_role = tf_m.paragraphs[0]
        p_m_role.text = "TEAM MEMBER"
        p_m_role.font.size = Pt(9)
        p_m_role.font.bold = True
        p_m_role.font.color.rgb = ACCENT_TEAL
        p_m_name = tf_m.add_paragraph()
        p_m_name.text = name
        p_m_name.font.size = Pt(14)
        p_m_name.font.bold = True
        p_m_name.font.color.rgb = TEXT_WHITE
        p_m_name.space_before = Pt(4)
        p_m_sub = tf_m.add_paragraph()
        p_m_sub.text = role
        p_m_sub.font.size = Pt(10)
        p_m_sub.font.color.rgb = TEXT_MUTED

    # =========================================================================
    # SLIDE 2: PROBLEM STATEMENT & MOTIVATION
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    add_header(s2, "The Growing Crisis of Digital Financial Fraud", "01 // PROBLEM STATEMENT", 2)

    # 3 Column Cards
    col_w = Inches(3.75)
    gap = Inches(0.24)
    y_pos = Inches(1.8)
    h_pos = Inches(5.1)

    # Col 1: Fraud Surge
    add_card(s2, Inches(0.8), y_pos, col_w, h_pos, "Massive Transaction Growth")
    tb2_1 = s2.shapes.add_textbox(Inches(1.0), y_pos + Inches(0.7), col_w - Inches(0.4), h_pos - Inches(0.9))
    tf2_1 = tb2_1.text_frame
    tf2_1.word_wrap = True
    bullets1 = [
        "Digital payments and UPI/card volumes are experiencing exponential global growth.",
        "Rapid transaction speeds (under 2 seconds) leave virtually zero window for manual review.",
        "Financial institutions lose tens of billions annually to unauthorized card-not-present fraud, identity theft, and account takeovers."
    ]
    for b in bullets1:
        p = tf2_1.add_paragraph()
        p.text = "• " + b
        p.font.size = Pt(13)
        p.font.color.rgb = TEXT_WHITE
        p.space_after = Pt(12)

    # Col 2: The Failure of Legacy Systems
    add_card(s2, Inches(0.8) + col_w + gap, y_pos, col_w, h_pos, "Limitations of Rule-Based Systems", border_color=ACCENT_RED)
    tb2_2 = s2.shapes.add_textbox(Inches(1.0) + col_w + gap, y_pos + Inches(0.7), col_w - Inches(0.4), h_pos - Inches(0.9))
    tf2_2 = tb2_2.text_frame
    tf2_2.word_wrap = True
    bullets2 = [
        "Static If-Else Rules: Fail to detect evolving, novel fraud attack vectors and money-mule networks.",
        "High False Positive Rates: Inconveniences legitimate customers by freezing genuine cards and transactions.",
        "Manual Operational Delay: Flagged queues become overwhelmed with backlog, delaying critical risk mitigation."
    ]
    for b in bullets2:
        p = tf2_2.add_paragraph()
        p.text = "• " + b
        p.font.size = Pt(13)
        p.font.color.rgb = TEXT_WHITE
        p.space_after = Pt(12)

    # Col 3: The AI Solution
    add_card(s2, Inches(0.8) + (col_w + gap) * 2, y_pos, col_w, h_pos, "Why Intelligent AI is Required", border_color=ACCENT_GREEN)
    tb2_3 = s2.shapes.add_textbox(Inches(1.0) + (col_w + gap) * 2, y_pos + Inches(0.7), col_w - Inches(0.4), h_pos - Inches(0.9))
    tf2_3 = tb2_3.text_frame
    tf2_3.word_wrap = True
    bullets3 = [
        "Behavioral Pattern Recognition: Analyzes non-linear correlations across amount, location, device, and velocity.",
        "Sub-Second Predictive Scoring: Yields risk probabilities and tiered recommendations instantly upon submission.",
        "Explainable Warning Signals: Tells human compliance officers exactly why a transaction was flagged."
    ]
    for b in bullets3:
        p = tf2_3.add_paragraph()
        p.text = "• " + b
        p.font.size = Pt(13)
        p.font.color.rgb = TEXT_WHITE
        p.space_after = Pt(12)

    # =========================================================================
    # SLIDE 3: PROJECT OBJECTIVES & KEY CAPABILITIES
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    add_header(s3, "Sentinel AI Core Objectives & Deliverables", "02 // OBJECTIVES & SCOPE", 3)

    obj_w = Inches(5.7)
    obj_h = Inches(2.35)
    
    # 4 Grid Cards
    cards_data = [
        ("Real-Time Transaction Risk Scoring", 
         "Develop an interactive analysis pipeline that computes fraud probability (0–100%) and categorizes risk into Low Risk, Review Required, or High Risk in under 100 milliseconds.",
         ACCENT_CYAN, Inches(0.8), Inches(1.8)),
        ("Explainable Intelligence & Decision Signals", 
         "Provide transparent warning signals and feature contributions explaining why an alert was triggered, rather than treating the ML model as an uninterpretable black box.",
         ACCENT_TEAL, Inches(6.8), Inches(1.8)),
        ("Full-Stack Investigation & Case Management", 
         "Build a dedicated Fraud Alerts Center allowing security analysts to triage alerts, update statuses (Open, Investigating, Resolved), and persist case notes directly to the database.",
         ACCENT_GREEN, Inches(0.8), Inches(4.4)),
        ("Comprehensive Visual Dashboard & Auditing", 
         "Deliver a SaaS-grade live dashboard with volume timeline analytics, risk distributions, paginated historical search, and CSV export for compliance auditing.",
         ACCENT_AMBER, Inches(6.8), Inches(4.4))
    ]

    for title, desc, col, x, y in cards_data:
        add_card(s3, x, y, obj_w, obj_h, title, border_color=col)
        tb = s3.shapes.add_textbox(x + Inches(0.2), y + Inches(0.65), obj_w - Inches(0.4), obj_h - Inches(0.8))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = desc
        p.font.size = Pt(13)
        p.font.color.rgb = TEXT_WHITE
        p.space_before = Pt(4)

    # =========================================================================
    # SLIDE 4: SYSTEM ARCHITECTURE & DATA FLOW
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    add_header(s4, "Multi-Tiered Full-Stack System Architecture", "03 // ARCHITECTURE & WORKFLOW", 4)

    # Left: Architecture Flow Layers
    add_card(s4, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.1), "Layered Architecture Breakdown")
    layers = [
        ("Client UI Layer (React 18 + Vite)", "Modern Single Page Application delivering high-performance UI, interactive forms, real-time charts, and toast notifications."),
        ("API Gateway Layer (FastAPI + Pydantic)", "High-performance asynchronous REST endpoints validating schema inputs, enforcing business logic, and routing calls."),
        ("ML Intelligence Engine (scikit-learn)", "Trained RandomForest model loading serialized artifacts to evaluate incoming feature vectors and generate probabilistic scores."),
        ("Data Persistence Layer (SQLite + SQLAlchemy)", "Relational persistence storing transactions, audit logs, alert histories, and case investigator notes across sessions.")
    ]
    tb4_l = s4.shapes.add_textbox(Inches(1.0), Inches(2.4), Inches(5.2), Inches(4.3))
    tf4_l = tb4_l.text_frame
    tf4_l.word_wrap = True
    for name, d in layers:
        p_h = tf4_l.add_paragraph()
        p_h.text = "▶ " + name
        p_h.font.size = Pt(12)
        p_h.font.bold = True
        p_h.font.color.rgb = ACCENT_CYAN
        p_d = tf4_l.add_paragraph()
        p_d.text = d
        p_d.font.size = Pt(11)
        p_d.font.color.rgb = TEXT_WHITE
        p_d.space_after = Pt(8)

    # Right: End-to-End Execution Pipeline (Step by step flow)
    add_card(s4, Inches(6.8), Inches(1.8), Inches(5.7), Inches(5.1), "End-to-End Processing Pipeline", border_color=ACCENT_TEAL)
    pipeline_steps = [
        ("Step 1: Transaction Input", "User submits transaction via web UI or API with amount, time, velocity, and device metrics."),
        ("Step 2: Server-Side Schema Validation", "FastAPI and Pydantic validate ranges, sanitize inputs, and prevent bad payloads."),
        ("Step 3: Feature Transformation", "Raw data is converted into numerical feature vectors matching model training specifications."),
        ("Step 4: Machine Learning Inference", "Random Forest classifier estimates probability of fraud based on learned patterns."),
        ("Step 5: Rule & Explanation Engine", "Identifies specific anomalies (e.g. nocturnal hour, velocity spike) to produce warnings."),
        ("Step 6: Database Persistence & Alerting", "Records saved in SQLite; transactions >= 65% risk automatically generate open alerts."),
        ("Step 7: Real-time UI Update", "Dashboard counters and alert feeds update immediately without requiring page refresh.")
    ]
    tb4_r = s4.shapes.add_textbox(Inches(7.0), Inches(2.4), Inches(5.3), Inches(4.3))
    tf4_r = tb4_r.text_frame
    tf4_r.word_wrap = True
    for st_title, st_desc in pipeline_steps:
        p_st = tf4_r.add_paragraph()
        p_st.text = st_title + " — " + st_desc
        p_st.font.size = Pt(11)
        p_st.font.color.rgb = TEXT_WHITE
        p_st.space_after = Pt(6)

    # =========================================================================
    # SLIDE 5: TECHNOLOGY STACK
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    add_header(s5, "Robust & Modern Technology Stack", "04 // TECH STACK", 5)

    stack_cols = [
        ("Frontend Client", [
            ("React 18 + TypeScript", "Component-based, type-safe UI architecture"),
            ("Vite Build Tool", "Lightning-fast HMR and bundle compilation"),
            ("Tailwind CSS", "Utility-first modern dark cybersecurity theme"),
            ("Recharts", "Interactive Area, Bar, and Donut visualizations"),
            ("Lucide React Icons", "Crisp vector icons across navigation & widgets"),
            ("React Router DOM", "Declarative client-side routing across 5 pages")
        ], ACCENT_CYAN),
        ("Backend Services", [
            ("Python 3.11+", "Clean, expressive environment for ML & backend"),
            ("FastAPI Framework", "Asynchronous, production-grade OpenAPI endpoints"),
            ("Pydantic v2", "Strict schema validation and error responses"),
            ("Uvicorn ASGI", "Ultra-fast ASGI web server for production"),
            ("CORS Middleware", "Restricted, secure cross-origin communication")
        ], ACCENT_TEAL),
        ("ML & Database", [
            ("scikit-learn", "RandomForestClassifier with ensemble bagging"),
            ("NumPy & pandas", "High-performance vector and tabular data manipulation"),
            ("joblib / Pickle", "Serialization of trained model artifacts"),
            ("SQLAlchemy ORM", "Declarative schema mapping and relation queries"),
            ("SQLite Database", "Serverless ACID persistence for auditability")
        ], ACCENT_GREEN)
    ]

    for i, (cat, items, col) in enumerate(stack_cols):
        x = Inches(0.8 + i * 4.0)
        add_card(s5, x, Inches(1.8), Inches(3.7), Inches(5.1), cat, border_color=col)
        tb = s5.shapes.add_textbox(x + Inches(0.2), Inches(2.4), Inches(3.3), Inches(4.3))
        tf = tb.text_frame
        tf.word_wrap = True
        for name, sub in items:
            p_n = tf.add_paragraph()
            p_n.text = "• " + name
            p_n.font.size = Pt(12)
            p_n.font.bold = True
            p_n.font.color.rgb = TEXT_WHITE
            p_s = tf.add_paragraph()
            p_s.text = "  " + sub
            p_s.font.size = Pt(10)
            p_s.font.color.rgb = TEXT_MUTED
            p_s.space_after = Pt(8)

    # =========================================================================
    # SLIDE 6: UI/UX DESIGN & BRAND IDENTITY
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    add_header(s6, "Visual Design & Cybersecurity Brand Identity", "05 // UI & USER EXPERIENCE", 6)

    # 4 Design Pillars
    pillars = [
        ("Dark Cyber Theme", "Custom dark navy palette (#0B132B) paired with deep slate cards, avoiding harsh eye fatigue during long security monitoring sessions.", ACCENT_CYAN, Inches(0.8), Inches(1.8)),
        ("Data-Dense Clean Layout", "Zero unnecessary clutter. Fine borders, restrained subtle glowing accents, and high-contrast typography designed for immediate cognitive clarity.", ACCENT_TEAL, Inches(6.8), Inches(1.8)),
        ("Visual Risk Indicators", "Intuitive three-tier color coding: Emerald Green for Low Risk, Golden Amber for Review Required, and Crimson Red for Critical Alerts.", ACCENT_GREEN, Inches(0.8), Inches(4.4)),
        ("Responsive & Dynamic", "Fully responsive desktop and tablet layouts with interactive drawer navigation, real-time loading skeletons, and helpful empty state displays.", ACCENT_AMBER, Inches(6.8), Inches(4.4))
    ]

    for title, desc, col, x, y in pillars:
        add_card(s6, x, y, Inches(5.7), Inches(2.35), title, border_color=col)
        tb = s6.shapes.add_textbox(x + Inches(0.2), y + Inches(0.65), Inches(5.3), Inches(1.5))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = desc
        p.font.size = Pt(13)
        p.font.color.rgb = TEXT_WHITE

    # =========================================================================
    # SLIDE 7: KEY FEATURE: TRANSACTION ANALYZER
    # =========================================================================
    s7 = prs.slides.add_slide(blank_layout)
    add_header(s7, "Core Feature: Interactive Transaction Analyzer", "06 // FEATURE SPOTLIGHT", 7)

    # Left: Input Parameters
    add_card(s7, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.1), "Multi-Factor Behavioral Input Form")
    form_fields = [
        ("Transaction Amount ($)", "Detects abnormally high charge deviations compared to historical spending baselines."),
        ("Transaction Hour (0-23)", "Analyzes temporal patterns; late-night/early-morning spikes trigger risk flags."),
        ("Merchant Category", "Differentiates risk levels (e.g., Electronics & Jewelry vs Groceries & Utilities)."),
        ("Transaction Type", "POS swipe, online payment, wire transfer, or ATM withdrawal."),
        ("Geographical Distance", "Distance (km) from customer home location; flags sudden physical anomalies."),
        ("Recent Transaction Velocity", "Number of transactions attempted within the last 24 hours."),
        ("Device Recognition", "Verifies whether the originating terminal or browser has been previously trusted.")
    ]
    tb7_l = s7.shapes.add_textbox(Inches(1.0), Inches(2.4), Inches(5.2), Inches(4.3))
    tf7_l = tb7_l.text_frame
    tf7_l.word_wrap = True
    for fn, fd in form_fields:
        p_fn = tf7_l.add_paragraph()
        p_fn.text = "• " + fn + ": "
        p_fn.font.size = Pt(11)
        p_fn.font.bold = True
        p_fn.font.color.rgb = ACCENT_CYAN
        p_fn.add_run().text = fd
        p_fn.font.color.rgb = TEXT_WHITE
        p_fn.space_after = Pt(4)

    # Right: Instant Result Card
    add_card(s7, Inches(6.8), Inches(1.8), Inches(5.7), Inches(5.1), "Instant Scoring & Explainability Engine", border_color=ACCENT_RED)
    tb7_r = s7.shapes.add_textbox(Inches(7.0), Inches(2.4), Inches(5.3), Inches(4.3))
    tf7_r = tb7_r.text_frame
    tf7_r.word_wrap = True
    
    res_items = [
        ("Sub-Second Fraud Probability", "Computes calibrated risk score (e.g., 88.4%) directly from model output."),
        ("Automated Tier Assignment", "Classified into High Risk, Review Required, or Low Risk based on preset thresholds."),
        ("Dynamic Warning Signals", "Explains specific suspicious behaviors detected: e.g., 'Unusual high velocity from unrecognized device at nocturnal hour'."),
        ("Audit Reference ID", "Generates unique tracking hash (e.g., TX-84920) stored in database for legal traceability."),
        ("Actionable Analyst Guidance", "Provides next steps: Approve, Freeze Card, or Route to Compliance Review Queue.")
    ]
    for rn, rd in res_items:
        p_rn = tf7_r.add_paragraph()
        p_rn.text = "▶ " + rn
        p_rn.font.size = Pt(12)
        p_rn.font.bold = True
        p_rn.font.color.rgb = ACCENT_TEAL
        p_rd = tf7_r.add_paragraph()
        p_rd.text = rd
        p_rd.font.size = Pt(11)
        p_rd.font.color.rgb = TEXT_WHITE
        p_rd.space_after = Pt(8)

    # =========================================================================
    # SLIDE 8: KEY FEATURE: LIVE OVERVIEW DASHBOARD
    # =========================================================================
    s8 = prs.slides.add_slide(blank_layout)
    add_header(s8, "Real-Time Transaction Monitoring Dashboard", "07 // FEATURE SPOTLIGHT", 8)

    # Top Row: 4 Metric Cards
    metrics = [
        ("Total Analyzed", "1,248", "All historical records", ACCENT_CYAN),
        ("High Risk Alerts", "42", "Transactions >= 65% risk", ACCENT_RED),
        ("Under Review", "89", "Requires human review", ACCENT_AMBER),
        ("Total Volume Monitored", "$248,910", "Aggregated transaction value", ACCENT_GREEN)
    ]
    for i, (m_title, m_val, m_sub, m_col) in enumerate(metrics):
        x = Inches(0.8 + i * 3.0)
        c = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(1.8), Inches(2.8), Inches(1.4))
        c.fill.solid()
        c.fill.fore_color.rgb = CARD_BG
        c.line.color.rgb = m_col
        c.line.width = Pt(1)
        tf = c.text_frame
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p0 = tf.paragraphs[0]
        p0.text = m_title.upper()
        p0.font.size = Pt(9)
        p0.font.bold = True
        p0.font.color.rgb = TEXT_MUTED
        p1 = tf.add_paragraph()
        p1.text = m_val
        p1.font.size = Pt(20)
        p1.font.bold = True
        p1.font.color.rgb = TEXT_WHITE
        p2 = tf.add_paragraph()
        p2.text = m_sub
        p2.font.size = Pt(9)
        p2.font.color.rgb = m_col

    # Bottom Left: Analytics Charts
    add_card(s8, Inches(0.8), Inches(3.5), Inches(6.0), Inches(3.4), "Visual Data Analytics (Recharts Integration)")
    tb8_l = s8.shapes.add_textbox(Inches(1.0), Inches(4.1), Inches(5.6), Inches(2.6))
    tf8_l = tb8_l.text_frame
    tf8_l.word_wrap = True
    charts_desc = [
        ("Temporal Activity Timeline", "Interactive Area Chart plotting transaction volume and suspicious frequency over time."),
        ("Risk Distribution Breakdown", "Donut/Pie charts displaying percentage proportions of Low, Review, and High Risk."),
        ("Zero Hardcoding Policy", "All chart coordinates and data points populate dynamically from backend `/api/dashboard/summary`.")
    ]
    for ct, cd in charts_desc:
        p = tf8_l.add_paragraph()
        p.text = "• " + ct + ": "
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = ACCENT_CYAN
        p.add_run().text = cd
        p.font.color.rgb = TEXT_WHITE
        p.space_after = Pt(8)

    # Bottom Right: Live Activity Stream
    add_card(s8, Inches(7.1), Inches(3.5), Inches(5.4), Inches(3.4), "Real-Time Transaction & Alert Feed", border_color=ACCENT_TEAL)
    tb8_r = s8.shapes.add_textbox(Inches(7.3), Inches(4.1), Inches(5.0), Inches(2.6))
    tf8_r = tb8_r.text_frame
    tf8_r.word_wrap = True
    feed_desc = [
        ("Instant Stream Updates", "Chronological list of recently processed transactions showing amount, category, and score."),
        ("Critical Threat Drawer", "Quick-action shortcuts allowing security officers to jump directly into high-risk alert details."),
        ("Graceful Empty States", "Intelligent visual placeholders guiding users to analyze transactions when database is newly initialized.")
    ]
    for ft, fd in feed_desc:
        p = tf8_r.add_paragraph()
        p.text = "▶ " + ft
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = ACCENT_TEAL
        p_sub = tf8_r.add_paragraph()
        p_sub.text = fd
        p_sub.font.size = Pt(10)
        p_sub.font.color.rgb = TEXT_WHITE
        p_sub.space_after = Pt(6)

    # =========================================================================
    # SLIDE 9: FRAUD ALERTS CENTER & INVESTIGATION WORKFLOW
    # =========================================================================
    s9 = prs.slides.add_slide(blank_layout)
    add_header(s9, "Fraud Alerts Center & Investigation Workflow", "08 // CASE MANAGEMENT", 9)

    add_card(s9, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.1), "Investigator Case Management Lifecycle")
    steps_case = [
        ("1. Automatic Alert Ingestion", "Any transaction scored above the critical risk threshold (e.g. 65%) is immediately placed into the Active Alert queue with an 'Open' status."),
        ("2. Multi-Criteria Triage", "Investigators can filter alerts by Risk Severity, Timestamp, Transaction Amount, or Status to prioritize high-value threats."),
        ("3. Investigation Status Transition", "Lifecycle tracking through 4 explicit states: [Open] ➜ [Investigating] ➜ [Resolved] or [False Positive]."),
        ("4. Persistent Audit Notes", "Security analysts record qualitative observations and evidence logs, saved permanently to SQLite with investigator timestamps.")
    ]
    tb9_l = s9.shapes.add_textbox(Inches(1.0), Inches(2.4), Inches(5.2), Inches(4.3))
    tf9_l = tb9_l.text_frame
    tf9_l.word_wrap = True
    for s_title, s_desc in steps_case:
        p_st = tf9_l.add_paragraph()
        p_st.text = s_title
        p_st.font.size = Pt(12)
        p_st.font.bold = True
        p_st.font.color.rgb = ACCENT_AMBER
        p_sd = tf9_l.add_paragraph()
        p_sd.text = s_desc
        p_sd.font.size = Pt(11)
        p_sd.font.color.rgb = TEXT_WHITE
        p_sd.space_after = Pt(8)

    # Right: Value to Banking / Enterprise Security
    add_card(s9, Inches(6.8), Inches(1.8), Inches(5.7), Inches(5.1), "Why Human-in-the-Loop AI Matters", border_color=ACCENT_CYAN)
    benefits = [
        ("Reduces False Positive Friction", "Allows operations teams to swiftly dismiss verified false alarms without disrupting user accounts permanently."),
        ("Continuous Feedback Loop", "Resolved investigations and marked false positives can serve as labeled ground truth for future retraining batches."),
        ("Compliance & Regulatory Readiness", "Maintains an immutable chain of custody with notes and status changes, satisfying KYC/AML auditing standards."),
        ("Prevents Catastrophic Financial Loss", "Enables rapid card freezing and merchant chargeback requests before fraudulent settlement completes.")
    ]
    tb9_r = s9.shapes.add_textbox(Inches(7.0), Inches(2.4), Inches(5.3), Inches(4.3))
    tf9_r = tb9_r.text_frame
    tf9_r.word_wrap = True
    for bt, bd in benefits:
        p_bt = tf9_r.add_paragraph()
        p_bt.text = "• " + bt
        p_bt.font.size = Pt(12)
        p_bt.font.bold = True
        p_bt.font.color.rgb = ACCENT_CYAN
        p_bd = tf9_r.add_paragraph()
        p_bd.text = bd
        p_bd.font.size = Pt(11)
        p_bd.font.color.rgb = TEXT_WHITE
        p_bd.space_after = Pt(8)

    # =========================================================================
    # SLIDE 10: TRANSACTION EXPLORER & COMPLIANCE REPORTING
    # =========================================================================
    s10 = prs.slides.add_slide(blank_layout)
    add_header(s10, "Transaction Explorer & Data Governance", "09 // SEARCH & COMPLIANCE", 10)

    col10_w = Inches(3.75)
    gap10 = Inches(0.24)

    # Card 1: Advanced Search
    add_card(s10, Inches(0.8), Inches(1.8), col10_w, Inches(5.1), "Searchable History Registry", border_color=ACCENT_CYAN)
    tb10_1 = s10.shapes.add_textbox(Inches(1.0), Inches(2.4), col10_w - Inches(0.4), Inches(4.3))
    tf10_1 = tb10_1.text_frame
    tf10_1.word_wrap = True
    items10_1 = [
        "Full-Text Search: Query across Transaction IDs, Merchant names, or Card references.",
        "Server-Side Pagination: Optimized response times handling thousands of stored rows.",
        "Multi-Tier Filtering: Filter simultaneously by Risk Score range, Amount interval, and Date bracket."
    ]
    for it in items10_1:
        p = tf10_1.add_paragraph()
        p.text = "▶ " + it
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_WHITE
        p.space_after = Pt(12)

    # Card 2: Deep Detail Inspection
    add_card(s10, Inches(0.8) + col10_w + gap10, Inches(1.8), col10_w, Inches(5.1), "Deep Transaction Inspection", border_color=ACCENT_TEAL)
    tb10_2 = s10.shapes.add_textbox(Inches(1.0) + col10_w + gap10, Inches(2.4), col10_w - Inches(0.4), Inches(4.3))
    tf10_2 = tb10_2.text_frame
    tf10_2.word_wrap = True
    items10_2 = [
        "Detailed Modal View: Inspect full behavioral payload without leaving table view.",
        "Device Fingerprint: Track IP, browser agent, and device recognition history.",
        "Geographic Coordinates: Visual comparison between client home origin and transaction terminal."
    ]
    for it in items10_2:
        p = tf10_2.add_paragraph()
        p.text = "▶ " + it
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_WHITE
        p.space_after = Pt(12)

    # Card 3: Secure CSV Export
    add_card(s10, Inches(0.8) + (col10_w + gap10) * 2, Inches(1.8), col10_w, Inches(5.1), "Secure CSV Export & Audits", border_color=ACCENT_GREEN)
    tb10_3 = s10.shapes.add_textbox(Inches(1.0) + (col10_w + gap10) * 2, Inches(2.4), col10_w - Inches(0.4), Inches(4.3))
    tf10_3 = tb10_3.text_frame
    tf10_3.word_wrap = True
    items10_3 = [
        "One-Click CSV Export: Export filtered transaction subsets for presentation and reports.",
        "Sanitized Data Export: Prevents formula injection and protects sensitive customer identifiers.",
        "College Demo Readiness: Allows live downloads of analyzed test sessions during academic viva."
    ]
    for it in items10_3:
        p = tf10_3.add_paragraph()
        p.text = "▶ " + it
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_WHITE
        p.space_after = Pt(12)

    # =========================================================================
    # SLIDE 11: MACHINE LEARNING PIPELINE & MODEL TRAINING
    # =========================================================================
    s11 = prs.slides.add_slide(blank_layout)
    add_header(s11, "Machine Learning Pipeline & Model Training", "10 // MACHINE LEARNING", 11)

    add_card(s11, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.1), "Random Forest Classifier Architecture")
    ml_points = [
        ("Algorithm Choice", "Random Forest Classifier utilizing ensemble bootstrap aggregating (bagging) over dozens of decision trees."),
        ("Non-Linear Detection", "Capable of identifying subtle, complex interactions between velocity, distance, and nocturnal transaction timing."),
        ("Overfitting Prevention", "Ensemble voting across decorrelated decision trees prevents memorization of noisy historical data."),
        ("Deterministic Fallback Demo", "Includes a synthetic data generator producing realistic fraud & legitimate distributions even without downloading heavy external CSV files.")
    ]
    tb11_l = s11.shapes.add_textbox(Inches(1.0), Inches(2.4), Inches(5.2), Inches(4.3))
    tf11_l = tb11_l.text_frame
    tf11_l.word_wrap = True
    for mn, md in ml_points:
        p_mn = tf11_l.add_paragraph()
        p_mn.text = "• " + mn
        p_mn.font.size = Pt(12)
        p_mn.font.bold = True
        p_mn.font.color.rgb = ACCENT_CYAN
        p_md = tf11_l.add_paragraph()
        p_md.text = md
        p_md.font.size = Pt(11)
        p_md.font.color.rgb = TEXT_WHITE
        p_md.space_after = Pt(6)

    # Right: Preprocessing & Training Pipeline
    add_card(s11, Inches(6.8), Inches(1.8), Inches(5.7), Inches(5.1), "Training Workflow & Real Dataset Adapter", border_color=ACCENT_TEAL)
    train_workflow = [
        ("Public Dataset Compatibility", "Built-in adapter maps public credit card fraud datasets (e.g., Kaggle European Credit Card Fraud) directly to the pipeline."),
        ("Stratified Data Split", "Strict 80/20 train/test split maintaining exact minority fraud class ratios across both sets to avoid training bias."),
        ("Target Leakage Prevention", "Scaler parameters and feature encodings are fitted strictly on training data, preventing test set leakage."),
        ("Artifact Serialization", "Model weights (`model.pkl`) and metadata (`model_metrics.json`) are stored cleanly in the `artifacts/` directory.")
    ]
    tb11_r = s11.shapes.add_textbox(Inches(7.0), Inches(2.4), Inches(5.3), Inches(4.3))
    tf11_r = tb11_r.text_frame
    tf11_r.word_wrap = True
    for wn, wd in train_workflow:
        p_wn = tf11_r.add_paragraph()
        p_wn.text = "▶ " + wn
        p_wn.font.size = Pt(12)
        p_wn.font.bold = True
        p_wn.font.color.rgb = ACCENT_TEAL
        p_wd = tf11_r.add_paragraph()
        p_wd.text = wd
        p_wd.font.size = Pt(11)
        p_wd.font.color.rgb = TEXT_WHITE
        p_wd.space_after = Pt(6)

    # =========================================================================
    # SLIDE 12: MODEL PERFORMANCE & EVALUATION METRICS
    # =========================================================================
    s12 = prs.slides.add_slide(blank_layout)
    add_header(s12, "Model Evaluation & Academic Fraud Metrics", "11 // METRICS & EVALUATION", 12)

    # 4 Grid metric cards
    w12 = Inches(5.7)
    h12 = Inches(2.35)

    eval_cards = [
        ("Why Accuracy is Insufficient in Fraud", 
         "In financial fraud datasets, 99.8% of transactions are legitimate and only 0.2% are fraud. A trivial dummy model predicting 'Always Legitimate' would achieve 99.8% accuracy while missing 100% of real-world fraud attacks.",
         ACCENT_RED, Inches(0.8), Inches(1.8)),
        ("Precision & Recall Balancing", 
         "• Precision: Out of all transactions flagged as fraud, how many were truly fraud? (Prevents customer harassment).\n• Recall (Sensitivity): Out of all actual fraud attacks, how many did we intercept? (Prevents monetary loss).",
         ACCENT_CYAN, Inches(6.8), Inches(1.8)),
        ("Confusion Matrix & F1-Score", 
         "Evaluates True Positives, False Positives, True Negatives, and False Negatives. The F1-Score harmonizes precision and recall into a single balanced metric suited for heavily imbalanced classes.",
         ACCENT_AMBER, Inches(0.8), Inches(4.4)),
        ("In-App Real-Time Evaluation Dashboard", 
         "Sentinel AI features a dedicated Model Performance page where users can inspect training timestamps, test sample counts, confusion matrix cells, and trigger model retraining with one click.",
         ACCENT_GREEN, Inches(6.8), Inches(4.4))
    ]

    for title, desc, col, x, y in eval_cards:
        add_card(s12, x, y, w12, h12, title, border_color=col)
        tb = s12.shapes.add_textbox(x + Inches(0.2), y + Inches(0.65), w12 - Inches(0.4), h12 - Inches(0.8))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = desc
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_WHITE

    # =========================================================================
    # SLIDE 13: CONCLUSION & FUTURE ENHANCEMENTS
    # =========================================================================
    s13 = prs.slides.add_slide(blank_layout)
    add_header(s13, "Project Summary & Future Roadmap", "12 // CONCLUSION & ROADMAP", 13)

    # Left: Project Achievements
    add_card(s13, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.1), "Key Project Achievements", border_color=ACCENT_GREEN)
    achievements = [
        ("Complete Full-Stack Engineering", "Delivered a fully integrated platform connecting React, FastAPI, SQLite, and scikit-learn without static mocks."),
        ("Real-Time Behavioral Scoring", "Instant multi-factor risk categorization (<35% Low, 35-64% Review, 65%+ High) with transparent warning signals."),
        ("Enterprise Investigation Workflow", "Complete lifecycle management (Open, Investigating, Resolved) with persistent notes."),
        ("Production Design Standards", "High-contrast dark cybersecurity aesthetics, responsive grids, and clean visual hierarchy.")
    ]
    tb13_l = s13.shapes.add_textbox(Inches(1.0), Inches(2.4), Inches(5.2), Inches(4.3))
    tf13_l = tb13_l.text_frame
    tf13_l.word_wrap = True
    for at, ad in achievements:
        p = tf13_l.add_paragraph()
        p.text = "✔ " + at
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = ACCENT_GREEN
        p_sub = tf13_l.add_paragraph()
        p_sub.text = ad
        p_sub.font.size = Pt(11)
        p_sub.font.color.rgb = TEXT_WHITE
        p_sub.space_after = Pt(8)

    # Right: Future Scope & Q&A
    add_card(s13, Inches(6.8), Inches(1.8), Inches(5.7), Inches(5.1), "Future Scope & Closing", border_color=ACCENT_CYAN)
    future_items = [
        ("Graph Neural Networks (GNNs)", "Identify complex coordinated fraud syndicates and money-mule rings across multi-hop transaction graphs."),
        ("Deep Anomaly Autoencoders", "Unsupervised deep learning for detecting zero-day fraud attacks with zero labeled training samples."),
        ("Enterprise Streaming Architecture", "Scale pipeline with Apache Kafka and Redis cluster for processing 100,000+ transactions per second."),
        ("Hardware Biometrics Integration", "Combine device telemetry with keystroke dynamics and behavioral biometric factors.")
    ]
    tb13_r = s13.shapes.add_textbox(Inches(7.0), Inches(2.4), Inches(5.3), Inches(3.2))
    tf13_r = tb13_r.text_frame
    tf13_r.word_wrap = True
    for ft, fd in future_items:
        p = tf13_r.add_paragraph()
        p.text = "✦ " + ft
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = ACCENT_CYAN
        p_sub = tf13_r.add_paragraph()
        p_sub.text = fd
        p_sub.font.size = Pt(10)
        p_sub.font.color.rgb = TEXT_WHITE
        p_sub.space_after = Pt(5)

    # Thank you / Q&A Banner
    qa_card = s13.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.0), Inches(5.4), Inches(5.3), Inches(1.3))
    qa_card.fill.solid()
    qa_card.fill.fore_color.rgb = RGBColor(14, 25, 50)
    qa_card.line.color.rgb = ACCENT_CYAN
    qa_card.line.width = Pt(1.5)
    tf_qa = qa_card.text_frame
    tf_qa.vertical_anchor = MSO_ANCHOR.MIDDLE
    p_qa1 = tf_qa.paragraphs[0]
    p_qa1.text = "Thank You! Questions & Discussion"
    p_qa1.font.size = Pt(15)
    p_qa1.font.bold = True
    p_qa1.font.color.rgb = TEXT_WHITE
    p_qa1.alignment = PP_ALIGN.CENTER
    p_qa2 = tf_qa.add_paragraph()
    p_qa2.text = "Presented by: Sanket Kushwaha (Team Lead), Anand Verma, Sandeep Mishra, Kunal"
    p_qa2.font.size = Pt(10)
    p_qa2.font.color.rgb = ACCENT_TEAL
    p_qa2.alignment = PP_ALIGN.CENTER
    p_qa2.space_before = Pt(4)

    # Save presentation
    output_path = "Sentinel_AI_Presentation.pptx"
    prs.save(output_path)
    print(f"Presentation successfully created at {output_path}")

if __name__ == "__main__":
    create_presentation()
